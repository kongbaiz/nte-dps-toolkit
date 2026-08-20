from docx import Document
from pathlib import Path
p=Path(r"D:\NTE_DPS_TOOL\.codex-fixtures\review-2026-08-19\ORIGINAL_REVIEW.docx")
out=Path(r"D:\NTE_DPS_TOOL\.codex-fixtures\review-2026-08-19\REVIEW_EXTRACT.txt")
doc=Document(p)
lines=[]
for i, para in enumerate(doc.paragraphs):
    text=para.text.strip()
    if text:
        lines.append(f"P{i:04d} [{para.style.name}] {text}")
for ti, table in enumerate(doc.tables):
    lines.append(f"\nTABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri,row in enumerate(table.rows):
        cells=[" ".join(c.text.split()) for c in row.cells]
        lines.append(f"T{ti}R{ri}: " + " | ".join(cells))
out.write_text("\n".join(lines), encoding="utf-8")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
print(out)
